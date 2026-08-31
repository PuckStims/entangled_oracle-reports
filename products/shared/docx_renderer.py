"""DOCX renderer for the shared semantic document model.

This module intentionally has no imports of report composers or report types.
"""
from __future__ import annotations

from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from .document_model import (
    CalloutBlock,
    CoverSection,
    DataTable,
    EventCard,
    Figure,
    Heading,
    KeyValueRow,
    MethodologyNote,
    MonthSection,
    NatalSummaryCard,
    PageBreak,
    Paragraph,
    PillRow,
    ProseSection,
    ReportDocument,
    ReportFooter,
    ReportHeader,
    Section,
    SupportCard,
)
from .document_styles import theme_for


class DocxRenderer:
    """Dispatches generic nodes to python-docx without product awareness."""

    def __init__(self, theme_id: str | None = None) -> None:
        self.theme = theme_for(theme_id)

    def render_to_path(self, report: ReportDocument, output_path: str | Path) -> Path:
        output = Path(output_path)
        output.parent.mkdir(parents=True, exist_ok=True)
        document = Document()
        self._configure_document(document)
        self._add_styles(document)
        for node in report.nodes:
            self.render_node(node, document)
        document.save(output)
        return output

    def _configure_document(self, document: Document) -> None:
        section = document.sections[0]
        # compact_reference_guide geometry with a named EO colour override.
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        section.header_distance = Inches(0.492)
        section.footer_distance = Inches(0.492)

    def _add_styles(self, document: Document) -> None:
        styles = document.styles

        def ensure(name: str, base: str = "Normal") -> Any:
            try:
                return styles[name]
            except KeyError:
                return styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)

        def style(name: str, size: float, *, bold: bool = False, color: str | None = None,
                  space_before: float = 0, space_after: float = 6, font: str | None = None) -> None:
            target = ensure(name)
            target.font.name = font or self.theme.body_font
            target.font.size = Pt(size)
            target.font.bold = bold
            target.font.color.rgb = RGBColor.from_string(color or self.theme.body_color)
            target.paragraph_format.space_before = Pt(space_before)
            target.paragraph_format.space_after = Pt(space_after)
            target.paragraph_format.line_spacing = 1.25

        style("EO Body", 11, space_after=6)
        style("EO Title", 30, bold=True, color=self.theme.accent_color, space_before=72, space_after=8, font=self.theme.heading_font)
        style("EO Heading 1", 16, bold=True, color=self.theme.accent_color, space_before=18, space_after=10, font=self.theme.heading_font)
        style("EO Heading 2", 13, bold=True, color=self.theme.accent_color, space_before=14, space_after=7, font=self.theme.heading_font)
        style("EO Heading 3", 12, bold=True, color=self.theme.body_color, space_before=10, space_after=5, font=self.theme.heading_font)
        style("EO Kicker", 8.5, bold=True, color=self.theme.muted_color, space_after=2)
        style("EO Callout", 11, color=self.theme.body_color, space_after=6)
        style("EO Event Title", 12, bold=True, color=self.theme.body_color, space_after=2, font=self.theme.heading_font)
        style("EO Event Metadata", 8.5, color=self.theme.muted_color, space_after=3)
        style("EO Caption", 8.5, color=self.theme.muted_color, space_after=8)
        style("EO Methodology", 9, color=self.theme.muted_color, space_before=8, space_after=8)
        style("EO Footer", 8, color=self.theme.muted_color, space_before=10, space_after=0)
        style("EO Table Header", 8.5, bold=True, color="FFFFFF", space_after=0)

    def render_node(self, node: Any, document: Document) -> None:
        handler = getattr(self, f"_render_{type(node).__name__.lower()}", None)
        if handler is None:
            raise TypeError(f"Unsupported document node: {type(node).__name__}")
        handler(node, document)

    def _render_coversection(self, node: CoverSection, document: Document) -> None:
        self._render_reportheader(node.header, document)
        for child in node.nodes:
            self.render_node(child, document)

    def _render_reportheader(self, node: ReportHeader, document: Document) -> None:
        title = document.add_paragraph(style="EO Title")
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title.add_run(node.title)
        if node.subtitle:
            subtitle = document.add_paragraph(style="EO Body")
            subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
            subtitle.add_run(node.subtitle)
        if node.metadata:
            self._render_key_value_table(node.metadata, document)

    def _render_section(self, node: Section, document: Document) -> None:
        if node.kicker:
            document.add_paragraph(node.kicker, style="EO Kicker")
        if node.title:
            document.add_paragraph(node.title, style="EO Heading 2")
        for child in node.nodes:
            self.render_node(child, document)

    def _render_monthsection(self, node: MonthSection, document: Document) -> None:
        if node.number is not None:
            document.add_paragraph(f"MONTH {node.number:02d}", style="EO Kicker")
        document.add_paragraph(node.title, style="EO Heading 2")
        for child in node.nodes:
            self.render_node(child, document)

    def _render_heading(self, node: Heading, document: Document) -> None:
        if node.kicker:
            document.add_paragraph(node.kicker, style="EO Kicker")
        document.add_paragraph(node.text, style=f"EO Heading {node.level}")

    def _render_paragraph(self, node: Paragraph, document: Document) -> None:
        paragraph = document.add_paragraph(style="EO Body")
        paragraph.add_run(node.text)

    def _render_prosesection(self, node: ProseSection, document: Document) -> None:
        document.add_paragraph(node.title, style="EO Heading 3")
        if node.subtitle:
            document.add_paragraph(node.subtitle, style="EO Event Metadata")
        self._render_paragraph(Paragraph(node.body, tone=node.tone), document)

    def _render_calloutblock(self, node: CalloutBlock, document: Document) -> None:
        table = document.add_table(rows=1, cols=1)
        self._set_table_geometry(table, (9360,))
        cell = table.cell(0, 0)
        self._set_cell_shading(cell, self.theme.background_color)
        self._set_cell_border(cell, self.theme.accent_for(node.tone))
        if node.title:
            cell.paragraphs[0].style = "EO Event Title"
            cell.paragraphs[0].add_run(node.title)
            paragraph = cell.add_paragraph(style="EO Callout")
        else:
            paragraph = cell.paragraphs[0]
            paragraph.style = "EO Callout"
        paragraph.add_run(node.body)

    def _render_eventcard(self, node: EventCard, document: Document) -> None:
        table = document.add_table(rows=1, cols=1)
        self._set_table_geometry(table, (9360,))
        cell = table.cell(0, 0)
        self._set_cell_border(cell, self.theme.accent_for(node.tone))
        cell.paragraphs[0].style = "EO Event Title"
        cell.paragraphs[0].add_run(node.title)
        if node.metadata:
            meta = cell.add_paragraph(style="EO Event Metadata")
            meta.add_run(" | ".join(f"{row.label}: {row.value}" for row in node.metadata if row.value))
        if node.body:
            cell.add_paragraph(node.body, style="EO Body")

    def _render_supportcard(self, node: SupportCard, document: Document) -> None:
        self._render_eventcard(EventCard(node.title, node.body, node.metadata, node.tone), document)

    def _render_natalsummarycard(self, node: NatalSummaryCard, document: Document) -> None:
        document.add_paragraph(node.title, style="EO Heading 3")
        self._render_key_value_table(node.rows, document)

    def _render_datatable(self, node: DataTable, document: Document) -> None:
        if node.caption:
            document.add_paragraph(node.caption, style="EO Caption")
        table = document.add_table(rows=1, cols=len(node.headers))
        table.style = "Table Grid"
        widths = self._column_widths(len(node.headers))
        self._set_table_geometry(table, widths)
        for index, header in enumerate(node.headers):
            cell = table.rows[0].cells[index]
            self._set_cell_shading(cell, self.theme.accent_color)
            cell.paragraphs[0].style = "EO Table Header"
            cell.paragraphs[0].add_run(header)
        for row in node.rows:
            cells = table.add_row().cells
            for index, value in enumerate(row):
                cells[index].text = value
                cells[index].paragraphs[0].style = "EO Body"
                self._set_cell_width(cells[index], widths[index])

    def _render_figure(self, node: Figure, document: Document) -> None:
        # python-docx cannot reliably add SVG in all Word versions.  Keep the
        # model source intact and provide an accessible, deterministic fallback.
        fallback = node.fallback_text or f"Figure unavailable in DOCX: {node.alt_text}"
        self._render_calloutblock(CalloutBlock(node.caption or "Figure", fallback, "structure"), document)
        if node.caption:
            document.add_paragraph(node.caption, style="EO Caption")

    def _render_pillrow(self, node: PillRow, document: Document) -> None:
        if node.items:
            document.add_paragraph(" | ".join(node.items), style="EO Event Metadata")

    def _render_methodologynote(self, node: MethodologyNote, document: Document) -> None:
        document.add_paragraph(node.text, style="EO Methodology")

    def _render_pagebreak(self, node: PageBreak, document: Document) -> None:
        document.add_page_break()

    def _render_reportfooter(self, node: ReportFooter, document: Document) -> None:
        paragraph = document.add_paragraph(node.text, style="EO Footer")
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def _render_key_value_table(self, rows: tuple[KeyValueRow, ...], document: Document) -> None:
        if not rows:
            return
        table = document.add_table(rows=1, cols=2)
        table.style = "Table Grid"
        widths = (2700, 6660)
        self._set_table_geometry(table, widths)
        for row_index, row in enumerate(rows):
            cells = table.rows[0].cells if row_index == 0 else table.add_row().cells
            cells[0].text = row.label
            cells[1].text = row.value
            for index, cell in enumerate(cells):
                cell.paragraphs[0].style = "EO Body"
                self._set_cell_width(cell, widths[index])

    @staticmethod
    def _column_widths(column_count: int) -> tuple[int, ...]:
        base, remainder = divmod(9360, column_count)
        return tuple(base + (1 if index < remainder else 0) for index in range(column_count))

    def _set_table_geometry(self, table: Any, widths: tuple[int, ...]) -> None:
        """Use fixed DXA geometry instead of Word's automatic table layout."""
        table.autofit = False
        tbl_pr = table._tbl.tblPr
        layout = tbl_pr.first_child_found_in("w:tblLayout")
        if layout is None:
            layout = OxmlElement("w:tblLayout")
        layout.set(qn("w:type"), "fixed")
        if layout.getparent() is None:
            tbl_pr.append(layout)
        table_width = tbl_pr.first_child_found_in("w:tblW")
        if table_width is None:
            table_width = OxmlElement("w:tblW")
        table_width.set(qn("w:w"), str(sum(widths)))
        table_width.set(qn("w:type"), "dxa")
        if table_width.getparent() is None:
            tbl_pr.append(table_width)
        indent = tbl_pr.first_child_found_in("w:tblInd")
        if indent is None:
            indent = OxmlElement("w:tblInd")
        indent.set(qn("w:w"), "120")
        indent.set(qn("w:type"), "dxa")
        if indent.getparent() is None:
            tbl_pr.append(indent)
        for grid_col, width in zip(table._tbl.tblGrid.gridCol_lst, widths):
            grid_col.set(qn("w:w"), str(width))
        for cell, width in zip(table.rows[0].cells, widths):
            self._set_cell_width(cell, width)

    @staticmethod
    def _set_cell_width(cell: Any, width: int) -> None:
        tc_pr = cell._tc.get_or_add_tcPr()
        tc_width = tc_pr.first_child_found_in("w:tcW")
        if tc_width is None:
            tc_width = OxmlElement("w:tcW")
        tc_width.set(qn("w:w"), str(width))
        tc_width.set(qn("w:type"), "dxa")
        if tc_width.getparent() is None:
            tc_pr.append(tc_width)
        margins = tc_pr.first_child_found_in("w:tcMar")
        if margins is None:
            margins = OxmlElement("w:tcMar")
        for edge, value in (("top", "80"), ("bottom", "80"), ("start", "120"), ("end", "120")):
            margin = margins.find(qn(f"w:{edge}"))
            if margin is None:
                margin = OxmlElement(f"w:{edge}")
                margins.append(margin)
            margin.set(qn("w:w"), value)
            margin.set(qn("w:type"), "dxa")
        if margins.getparent() is None:
            tc_pr.append(margins)

    @staticmethod
    def _set_cell_shading(cell: Any, fill: str) -> None:
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), fill)
        cell._tc.get_or_add_tcPr().append(shading)

    @staticmethod
    def _set_cell_border(cell: Any, color: str) -> None:
        tc_pr = cell._tc.get_or_add_tcPr()
        borders = tc_pr.first_child_found_in("w:tcBorders")
        if borders is None:
            borders = OxmlElement("w:tcBorders")
            tc_pr.append(borders)
        left = OxmlElement("w:left")
        left.set(qn("w:val"), "single")
        left.set(qn("w:sz"), "16")
        left.set(qn("w:color"), color)
        borders.append(left)
