import base64

from docx import Document

from products.shared.docx_renderer import DocxRenderer
from products.shared.document_model import (
    CalloutBlock,
    DataTable,
    EventCard,
    Figure,
    Heading,
    KeyValueRow,
    MethodologyNote,
    PageBreak,
    Paragraph,
    ReportDocument,
    ReportFooter,
    Section,
)


def test_docx_renderer_writes_reopenable_document_with_semantic_content(tmp_path):
    output = tmp_path / "synthetic.docx"
    report = ReportDocument(
        title="Synthetic",
        theme_id="vibrant",
        nodes=(
            Section(
                "Overview",
                (
                    Heading("Overview", level=1),
                    Paragraph("Body text."),
                    CalloutBlock("Important", "A semantic callout.", tone="threshold"),
                    EventCard("Event title", "Event body.", (KeyValueRow("Date", "Jan 1"),), tone="pressure"),
                    DataTable(("Key", "Value"), (("One", "1"),)),
                    Figure("image/svg+xml", "<svg/>", "Example figure", fallback_text="Figure fallback text."),
                    PageBreak(),
                    MethodologyNote("Methodology text."),
                    ReportFooter("Footer text."),
                ),
                marker="overview",
            ),
        ),
    )

    rendered = DocxRenderer(theme_id=report.theme_id).render_to_path(report, output)
    reopened = Document(rendered)
    paragraph_text = [paragraph.text for paragraph in reopened.paragraphs]
    style_names = {style.name for style in reopened.styles}

    assert rendered.exists()
    assert "Overview" in paragraph_text
    assert "Body text." in paragraph_text
    assert "Methodology text." in paragraph_text
    assert "Footer text." in paragraph_text
    assert "EO Heading 1" in style_names
    assert "EO Event Title" in style_names
    assert len(reopened.tables) >= 3
    assert any("Figure fallback text." in cell.text for table in reopened.tables for row in table.rows for cell in row.cells)
    first_table_width = reopened.tables[0]._tbl.tblPr.first_child_found_in("w:tblW")
    assert first_table_width.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}w") == "9360"


def test_renderer_source_does_not_know_report_types():
    source = ("products/shared/docx_renderer.py").replace("\\", "/")
    contents = open(source, encoding="utf-8").read()

    for report_type in ("year_ahead", "synastry", "horoscope", "personal_forecast"):
        assert report_type not in contents


def test_renderer_embeds_raster_figure_with_accessible_description(tmp_path):
    # Valid one-pixel PNG; keeping it inline avoids a graphics test dependency.
    png = base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVQIHWP4z8DwHwAFgAI/ScL0NwAAAABJRU5ErkJggg==")
    report = ReportDocument("Raster figure", (Figure("image/png", png, "One-pixel test image", caption="Raster caption"),))
    output = tmp_path / "raster.docx"

    DocxRenderer().render_to_path(report, output)
    reopened = Document(output)

    assert len(reopened.inline_shapes) == 1
    assert "Raster caption" in "\n".join(paragraph.text for paragraph in reopened.paragraphs)


def test_renderer_converts_svg_figure_to_embedded_docx_image(tmp_path):
    svg = '<svg xmlns="http://www.w3.org/2000/svg" width="24" height="24"><circle cx="12" cy="12" r="10" fill="#6c4a8a"/></svg>'
    report = ReportDocument("SVG figure", (Figure("image/svg+xml", svg, "Purple circle", caption="SVG caption", width_inches=1.0),))
    output = tmp_path / "svg.docx"

    DocxRenderer().render_to_path(report, output)
    reopened = Document(output)

    assert len(reopened.inline_shapes) == 1
    assert "SVG caption" in "\n".join(paragraph.text for paragraph in reopened.paragraphs)
